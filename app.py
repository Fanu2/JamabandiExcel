import streamlit as st
import camelot
import pandas as pd
from docx import Document
import io

def extract_tables_camelot(file):
    tables = camelot.read_pdf(file, pages='all', flavor='lattice')
    if tables.n > 1:
        df_list = [t.df for t in tables]
        df = pd.concat(df_list, ignore_index=True)
    else:
        df = tables[0].df
    df.columns = df.iloc[0]
    df = df[1:].reset_index(drop=True)
    return df

def table_to_docx(df, filename="table.docx"):
    doc = Document()
    table = doc.add_table(rows=df.shape[0]+1, cols=df.shape[1])
    for j, column in enumerate(df.columns):
        table.cell(0, j).text = str(column)
    for i in range(df.shape[0]):
        for j in range(df.shape[1]):
            table.cell(i+1, j).text = str(df.iloc[i, j])
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf

st.title("Jamabandi Table Extractor (Camelot)")

uploaded_file = st.file_uploader("Upload PDF", type=["pdf"])

if uploaded_file is not None:
    df = extract_tables_camelot(uploaded_file)
    st.success("Full table extracted successfully!")
    st.write("Preview:")
    st.dataframe(df)

    st.download_button(
        "Download CSV",
        df.to_csv(index=False).encode('utf-8'),
        file_name="jamabandi_table.csv",
        mime="text/csv"
    )

    excel_buffer = io.BytesIO()
    df.to_excel(excel_buffer, index=False)
    excel_buffer.seek(0)
    st.download_button(
        "Download Excel",
        excel_buffer.getvalue(),
        file_name="jamabandi_table.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

    docx_buffer = table_to_docx(df)
    st.download_button(
        "Download DOCX",
        docx_buffer.getvalue(),
        file_name="jamabandi_table.docx",
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )